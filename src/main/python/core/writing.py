# Author: Dahir Muhammad Dahir
# Date: 18th-05-2019 05:39 PM
# this module requires python3

import openpyxl as excel
from core.error import Error


class WriteExcel:
    def __init__(self, tables=None, table_fields=None, filename=None):
        self.tables = tables
        self.table_fields = table_fields
        self.output_filename = filename
        self.workbook = excel.Workbook()
        self.worksheet = self.workbook.active

    def write_heading(self):
        try:
            self.worksheet.append(self.table_fields)
        except Exception as e:
            Error.error_handler(e)

    def write_ending(self):
        try:
            space = []
            for _ in self.table_fields:
                space.append(" ")

            for i in range(3):
                self.worksheet.append(space)
        except Exception as e:
            Error.error_handler(e)

    def write_all(self):
        try:
            self.write_heading()
            for table in self.tables:
                self.worksheet.append(table)
            self.write_ending()
            self.workbook.save(self.output_filename)
        except Exception as e:
            Error.error_handler(e)

    def write_daily(self):
        try:
            current_day = self.tables[0]["B"][:11]
            self.write_heading()
            for table in self.tables:
                if current_day == table["B"][:11]:
                    self.worksheet.append(table)
                else:
                    self.write_ending()
                    current_day = table["B"][:11]
                    self.write_heading()
                    self.worksheet.append(table)
            self.write_ending()
            self.workbook.save(self.output_filename)
        except Exception as e:
            Error.error_handler(e)

    def write_monthly(self):
        try:
            current_month = self.tables[0]["B"][:8]
            self.write_heading()
            for table in self.tables:
                if current_month == table["B"][:8]:
                    self.worksheet.append(table)
                else:
                    self.write_ending()
                    current_month = table["B"][:8]
                    self.write_heading()
                    self.worksheet.append(table)
            self.write_ending()
            self.workbook.save(self.output_filename)
        except Exception as e:
            Error.error_handler(e)
